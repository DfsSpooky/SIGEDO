import os
import random
import re

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.models.academic import Especialidad
from core.utils.docente_identity import normalize_docente_key

User = get_user_model()


class Command(BaseCommand):
    help = "Registra docentes desde el Excel de malla curricular (sin crear cursos)."

    def add_arguments(self, parser):
        parser.add_argument(
            "--archivo",
            type=str,
            default="HORARIO DE CLASES 2026-A.xlsx",
            help="Nombre de archivo dentro de gestion_docentes/data_temp/ o ruta absoluta.",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        hojas_especialidades = {
            "M-F": "MATEMATICA - FISICA",
            "TIT": "TECNOLOGIA INFORMATICA Y TELECOMUNICACIONES",
            "BBQQ": "BIOLOGIA Y QUIMICA",
            "CL": "COMUNICACIÓN Y LITERATURA",
            "IDIOMAS": "LENGUAS EXTRANJERAS: INGLÉS FRANCÉS",
            "HISTORIA": "HISTORIA CCSS Y TURISMO",
            "FILOSOFÍA": "CCSS FILOSOFÍA Y PSICOLOGÍA EDUCATIVA",
        }

        regex_separador = re.compile(
            r"^(.*?)\s+(Dr\.|Dra\.|Mg\.|Contrata|Contrato|Alberto Cabrera|Lilia Matos)(.*)$",
            re.IGNORECASE,
        )

        archivo = options["archivo"]
        ruta_excel = archivo
        if not os.path.isabs(archivo):
            candidatos = [
                os.path.join(os.getcwd(), "data_temp", archivo),
                os.path.join(os.getcwd(), "gestion_docentes", "data_temp", archivo),
            ]
            ruta_excel = next((p for p in candidatos if os.path.exists(p)), candidatos[0])

        if not os.path.exists(ruta_excel):
            raise CommandError(
                f"No se encontró el archivo Excel. Rutas probadas: {ruta_excel}"
            )

        self.stdout.write(self.style.SUCCESS(f"Leyendo Excel: {ruta_excel}"))
        wb = openpyxl.load_workbook(ruta_excel, data_only=True)

        docentes_creados = 0
        docentes_actualizados = 0
        placeholders_omitidos = 0
        docentes_cache = {}

        for nombre_hoja, nombre_especialidad in hojas_especialidades.items():
            if nombre_hoja not in wb.sheetnames:
                self.stdout.write(
                    self.style.WARNING(f'Omitiendo: no se encontró la pestaña "{nombre_hoja}"')
                )
                continue

            sheet = wb[nombre_hoja]
            especialidad = Especialidad.objects.filter(nombre=nombre_especialidad).first()
            if not especialidad:
                self.stdout.write(
                    self.style.WARNING(
                        f'Especialidad "{nombre_especialidad}" no existe. Ejecuta primero "configurar_institucion".'
                    )
                )
                continue

            self.stdout.write(f"Procesando hoja {nombre_hoja} -> {nombre_especialidad}")

            for row in sheet.iter_rows(values_only=True):
                fila = [str(celda).strip() if celda is not None else "" for celda in row]
                if len(fila) < 8:
                    continue

                for celda in fila[3:8]:
                    celda = " ".join(celda.split()).strip()
                    if not celda:
                        continue

                    # Saltar cabeceras/horas/días.
                    lower = celda.lower()
                    if lower in {
                        "lunes",
                        "martes",
                        "miercoles",
                        "miércoles",
                        "jueves",
                        "viernes",
                        "sábado",
                        "sabado",
                        "domingo",
                        "hora",
                    }:
                        continue

                    docente_full_name = "Sin Asignar"
                    nombre_docente = "Sin Asignar"
                    match = regex_separador.match(celda)
                    if match:
                        titulo = match.group(2).strip()
                        nombre_docente = match.group(3).strip()
                        docente_full_name = f"{titulo} {nombre_docente}".strip()
                        if titulo.lower() in {"alberto cabrera", "lilia matos"}:
                            nombre_docente = docente_full_name
                    else:
                        # Celda sin separador claro: podría ser curso sin docente.
                        continue

                    if "contrat" in docente_full_name.lower() or "sin asignar" in docente_full_name.lower():
                        placeholders_omitidos += 1
                        continue

                    clean_name = normalize_docente_key(nombre_docente)
                    if not clean_name:
                        continue

                    base_username = clean_name[:30]
                    email = f"{clean_name}@undac.edu.pe"
                    docente = docentes_cache.get(clean_name)

                    if not docente:
                        # 1) Reutiliza usuario canónico (evita duplicados por variaciones mínimas de nombre/título).
                        docente = User.objects.filter(username=base_username).first()

                    if not docente:
                        # 2) Si ya existe una variante con sufijo (ej. juanperez1), reutilizarla.
                        candidatos = User.objects.filter(
                            username__startswith=base_username,
                            is_staff=False,
                            is_superuser=False,
                        ).order_by("id")
                        for candidato in candidatos:
                            nombre_guardado = " ".join(
                                p for p in [candidato.first_name, candidato.last_name] if p
                            ).strip()
                            if (
                                normalize_docente_key(nombre_guardado) == clean_name
                                or normalize_docente_key(candidato.first_name) == clean_name
                            ):
                                docente = candidato
                                break

                    if not docente:
                        # 3) Fallback por email canónico.
                        docente = User.objects.filter(email__iexact=email).first()

                    if not docente:
                        username = base_username
                        idx = 1
                        while User.objects.filter(username=username).exists():
                            suffix = str(idx)
                            username = f"{base_username[: max(1, 30 - len(suffix))]}{suffix}"
                            idx += 1

                        fake_dni = str(random.randint(10000000, 99999999))
                        while User.objects.filter(dni=fake_dni).exists():
                            fake_dni = str(random.randint(10000000, 99999999))

                        docente = User.objects.create(
                            username=username,
                            email=email,
                            first_name=docente_full_name[:150],
                            dni=fake_dni,
                            is_staff=False,
                            is_superuser=False,
                        )
                        docentes_creados += 1
                    else:
                        updated = False
                        if not docente.first_name:
                            docente.first_name = docente_full_name[:150]
                            updated = True
                        if not docente.email:
                            docente.email = email
                            updated = True
                        if updated:
                            docente.save(update_fields=["first_name", "email"])
                            docentes_actualizados += 1

                    docentes_cache[clean_name] = docente
                    docente.especialidades.add(especialidad)

        self.stdout.write(self.style.SUCCESS("\n===================================="))
        self.stdout.write(self.style.SUCCESS("¡REGISTRO DE DOCENTES COMPLETADO!"))
        self.stdout.write(self.style.SUCCESS(f"Docentes creados: {docentes_creados}"))
        self.stdout.write(self.style.SUCCESS(f"Docentes actualizados: {docentes_actualizados}"))
        self.stdout.write(self.style.SUCCESS(f"Placeholders omitidos: {placeholders_omitidos}"))
        self.stdout.write(self.style.SUCCESS("====================================\n"))
