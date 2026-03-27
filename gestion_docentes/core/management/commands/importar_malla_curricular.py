import re
import os
import random
import openpyxl
from datetime import time, date
from django.core.management.base import BaseCommand
from core.models.academic import Especialidad, Curso, Carrera, Semestre
from core.models.scheduling import FranjaHoraria
from core.models.users import Docente
from core.utils.docente_identity import normalize_docente_key

class Command(BaseCommand):
    help = 'Importa Especialidades, Docentes y Cursos directamente desde el archivo Excel original.'

    def add_arguments(self, parser):
        parser.add_argument(
            "--archivo",
            type=str,
            default="HORARIO DE CLASES 2026-A.xlsx",
            help="Nombre de archivo dentro de data_temp/ o ruta absoluta.",
        )
        parser.add_argument(
            "--reset",
            action="store_true",
            help="Borra cursos y docentes no admin antes de importar.",
        )

    def handle(self, *args, **kwargs):
        # Mapeo de los nombres de las PESTAÑAS de tu Excel a las Especialidades reales
        hojas_especialidades = {
            'M-F': 'MATEMATICA - FISICA',
            'TIT': 'TECNOLOGIA INFORMATICA Y TELECOMUNICACIONES',
            'BBQQ': 'BIOLOGIA Y QUIMICA',
            'CL': 'COMUNICACIÓN Y LITERATURA',
            'IDIOMAS': 'LENGUAS EXTRANJERAS: INGLÉS FRANCÉS',
            'HISTORIA': 'HISTORIA CCSS Y TURISMO',
            'FILOSOFÍA': 'CCSS FILOSOFÍA Y PSICOLOGÍA EDUCATIVA'
        }

        # Expresión regular para separar Curso de Docente
        regex_separador = re.compile(
            r'^(.*?)\s+(Dr\.|Dra\.|Mg\.|Contrata|Contrato|Alberto Cabrera|Lilia Matos)(.*)$', 
            re.IGNORECASE
        )

        archivo = kwargs["archivo"]
        ruta_excel = archivo
        if not os.path.isabs(archivo):
            candidatos = [
                os.path.join(os.getcwd(), "data_temp", archivo),
                os.path.join(os.getcwd(), "gestion_docentes", "data_temp", archivo),
            ]
            ruta_excel = next((p for p in candidatos if os.path.exists(p)), candidatos[0])

        if not os.path.exists(ruta_excel):
            self.stdout.write(self.style.ERROR(f'No se encontró el archivo Excel en: {ruta_excel}'))
            return

        if kwargs["reset"]:
            # Modo explícito destructivo, nunca por defecto en producción.
            self.stdout.write(self.style.WARNING('Modo --reset activo: limpiando Cursos y Docentes no admin...'))
            Curso.objects.all().delete()
            Docente.objects.filter(is_staff=False, is_superuser=False).delete()
        else:
            self.stdout.write(self.style.WARNING('Modo seguro: sin borrado previo. Usa --reset solo si realmente quieres reiniciar.'))

        # Abrir el Excel directamente
        self.stdout.write(self.style.SUCCESS(f'Leyendo el archivo Excel...'))
        wb = openpyxl.load_workbook(ruta_excel, data_only=True)
        
        # 1. Asegurar Semestre 2026-A Activo
        semestre_2026a, s_created = Semestre.objects.get_or_create(
            nombre="2026-A",
            defaults={
                'tipo': 'IMPAR',
                'estado': 'ACTIVO',
                'fecha_inicio': date(2026, 3, 1),
                'fecha_fin': date(2026, 7, 31)
            }
        )
        if not s_created and semestre_2026a.estado != 'ACTIVO':
            semestre_2026a.estado = 'ACTIVO'
            semestre_2026a.save()
        self.stdout.write(self.style.SUCCESS(f'Semestre {semestre_2026a.nombre} asegurado como ACTIVO.'))

        # 2. Asegurar Franjas Horarias 
        if getattr(FranjaHoraria.objects, 'count')() == 0:
            self.stdout.write(self.style.WARNING(f'Creando Franjas Horarias por defecto (Mañana y Tarde)...'))
            franjas_manana = [
                (time(7, 0), time(7, 50)), (time(7, 50), time(8, 40)), (time(8, 40), time(9, 30)),
                (time(9, 30), time(10, 20)), (time(10, 20), time(11, 10)), (time(11, 10), time(12, 0)),
                (time(12, 0), time(12, 50)), (time(12, 50), time(13, 40)),
            ]
            for inc, fin in franjas_manana:
                FranjaHoraria.objects.get_or_create(turno="MANANA", hora_inicio=inc, hora_fin=fin)
                
            franjas_tarde = [
                (time(14, 0), time(14, 50)), (time(14, 50), time(15, 40)), (time(15, 40), time(16, 30)),
                (time(16, 30), time(17, 20)), (time(17, 20), time(18, 10)), (time(18, 10), time(19, 0)),
                (time(19, 0), time(19, 50)), (time(19, 50), time(20, 40)),
            ]
            for inc, fin in franjas_tarde:
                FranjaHoraria.objects.get_or_create(turno="TARDE", hora_inicio=inc, hora_fin=fin)
            self.stdout.write(self.style.SUCCESS('Franjas de Mañana y Tarde creadas exitosamente.'))
        
        docentes_creados = 0
        cursos_creados = 0

        for nombre_hoja, nombre_especialidad in hojas_especialidades.items():
            if nombre_hoja not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f'Omitiendo: No se encontró la pestaña "{nombre_hoja}"'))
                continue
            
            sheet = wb[nombre_hoja]
            especialidad, _ = Especialidad.objects.get_or_create(nombre=nombre_especialidad)
            carrera_obj, _ = Carrera.objects.get_or_create(nombre="ESCUELA PROFESIONAL DE EDUCACIÓN SECUNDARIA")
            self.stdout.write(self.style.SUCCESS(f'Procesando pestaña: {nombre_hoja} -> {especialidad.nombre}...'))

            semestre_actual = "I" # Por defecto
            roman_to_int = {
                'I': 1, 'II': 2, 'III': 3, 'IV': 4, 'V': 5,
                'VI': 6, 'VII': 7, 'VIII': 8, 'IX': 9, 'X': 10
            }
            
            for row in sheet.iter_rows(values_only=True):
                fila = [str(celda).strip() if celda is not None else '' for celda in row]
                
                if len(fila) < 4:
                    continue
                
                col0 = fila[0].upper()
                if col0 in ['I', 'III', 'V', 'VII', 'IX']:
                    semestre_actual = col0

                for celda in fila[3:8]:
                    celda = " ".join(celda.split()).strip() 
                    if not celda or celda == '':
                        continue
                    
                    match = regex_separador.match(celda)
                    if match:
                        nombre_curso = match.group(1).strip()
                        titulo = match.group(2).strip()
                        nombre_docente = match.group(3).strip()
                        docente_full_name = f"{titulo} {nombre_docente}".strip()
                        
                        # Fix for hardcoded names in the regex as 'titulos'
                        if titulo.lower() in ["alberto cabrera", "lilia matos"]:
                            nombre_docente = docente_full_name
                        
                        if "contrat" in docente_full_name.lower():
                            docente_full_name = "Docente Por Contratar"
                            nombre_docente = "Docente Por Contratar"
                    else:
                        nombre_curso = celda
                        docente_full_name = "Sin Asignar"
                        nombre_docente = "Sin Asignar"

                    # Saltar si el "curso" es solo un dia de la semana (cabeceras del excel)
                    dias_semana_nombres = ['lunes', 'martes', 'miércoles', 'miercoles', 'jueves', 'viernes', 'sabado', 'sábado', 'domingo']
                    if nombre_curso.lower() in dias_semana_nombres:
                        continue

                    clean_name = normalize_docente_key(nombre_docente)
                    
                    if "contrat" in clean_name or "sinasignar" in clean_name:
                        clean_name = "docenteporcontratar"
                        docente_full_name = "Docente Por Contratar"

                    email_base = f"{clean_name}@undac.edu.pe"
                    username_base = clean_name[:30]

                    # 1. Buscamos por username canónico.
                    docente_obj = Docente.objects.filter(username=username_base).first()
                    if not docente_obj:
                        # 2. Reutilizar variantes existentes con sufijo si representan al mismo docente.
                        candidatos = Docente.objects.filter(
                            username__startswith=username_base,
                            is_staff=False,
                            is_superuser=False,
                        ).order_by("id")
                        for candidato in candidatos:
                            nombre_guardado = " ".join(
                                p for p in [candidato.first_name, candidato.last_name] if p
                            ).strip()
                            if normalize_docente_key(nombre_guardado) == clean_name:
                                docente_obj = candidato
                                break
                    if not docente_obj:
                        # 3. Fallback por email canónico.
                        docente_obj = Docente.objects.filter(email__iexact=email_base).first()
                    
                    if not docente_obj:
                        # Necesitamos generar un DNI temporal único porque tu modelo lo exige
                        fake_dni = str(random.randint(10000000, 99999999))
                        while Docente.objects.filter(dni=fake_dni).exists():
                            fake_dni = str(random.randint(10000000, 99999999))
                            
                        docente_obj = Docente.objects.create(
                            username=username_base,
                            email=email_base,
                            first_name=docente_full_name[:150],
                            dni=fake_dni
                        )
                        docentes_creados += 1

                    # 2. Crear el Curso
                    semestre_cursado_val = roman_to_int.get(semestre_actual.upper(), 1)
                    curso_obj, c_created = Curso.objects.get_or_create(
                        nombre=nombre_curso,
                        carrera=carrera_obj,
                        semestre_cursado=semestre_cursado_val,
                        defaults={
                            'docente': docente_obj,
                            'tipo_curso': 'ESPECIALIDAD',
                            'semestre': semestre_2026a
                        }
                    )
                    
                    # Lo vinculamos a la especialidad
                    curso_obj.especialidades.add(especialidad)
                    
                    if c_created: cursos_creados += 1

        self.stdout.write(self.style.SUCCESS(f'\n===================================='))
        self.stdout.write(self.style.SUCCESS(f'¡IMPORTACIÓN DESDE EXCEL COMPLETADA!'))
        self.stdout.write(self.style.SUCCESS(f'Nuevos Docentes registrados: {docentes_creados}'))
        self.stdout.write(self.style.SUCCESS(f'Cursos únicos registrados: {cursos_creados}'))
        self.stdout.write(self.style.SUCCESS(f'====================================\n'))
